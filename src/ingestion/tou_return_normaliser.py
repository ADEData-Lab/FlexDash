"""
Non-template TOU return normalisation (Option 1).

Some contributors provide Time-of-Use (TOU) / implicit flexibility data in a spreadsheet that
does not follow the ADE Data Template V1.2 layout.

This module converts such a return into a V1.2-shaped workbook so that:
  - administrators can open/review it alongside other returns, and
  - the ingest pipeline can parse it consistently.

Important:
  - The original submission is NOT modified.
  - The normalised workbook is an interpretation/mapping layer and must be treated with
    appropriate caveats.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Callable

import pandas as pd
from openpyxl import load_workbook


@dataclass(frozen=True)
class TOUReturnNormalisationResult:
    output_path: Path
    notes: List[str]


def _cell_text(v: Any) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def _row_has_all(row: List[Any], *needles: str) -> bool:
    t = " ".join(_cell_text(v).lower() for v in row if _cell_text(v))
    return all(n.lower() in t for n in needles)


def _find_row(df: pd.DataFrame, predicate) -> Optional[int]:
    for i in range(len(df)):
        if predicate(df.iloc[i].tolist()):
            return i
    return None


def _default_baseline_matcher(comparison: str) -> bool:
    return "flexible" in (comparison or "").strip().lower()


def _extract_tou_blocks(
    raw_path: Path,
    *,
    baseline_matcher: Callable[[str], bool] = _default_baseline_matcher,
) -> Dict[str, Any]:
    """
    Extract the blocks we can reliably map into the V1.2 template:
      - Asset counts at dates (used for template Part 2)
      - Tariff metadata (template Part 4.1)
      - Total turn-up / turn-down differences (template Part 4.2)
      - Daily average busiest HH MWh + MW (template Part 4.3)

    The raw sheet may contain multiple comparison baselines. We select rows whose
    'comparison' value matches `baseline_matcher`.
    """
    df = pd.read_excel(raw_path, header=None, sheet_name=0)

    # ---- Block 1: asset counts (dates) ----
    counts_header = _find_row(
        df,
        lambda r: _row_has_all(r, "assets")
        and _row_has_all(r, "1st nov 2024")
        and any(
            hasattr(v, "year") and getattr(v, "year", None) == 2025 and getattr(v, "month", None) == 2
            for v in r
        ),
    )
    if counts_header is None:
        raise ValueError("TOU normaliser: could not locate the asset counts header row")

    counts_cols = df.iloc[counts_header].tolist()
    feb_col = None
    nov_col = None
    for j, v in enumerate(counts_cols):
        if hasattr(v, "year") and getattr(v, "year", None) == 2025 and getattr(v, "month", None) == 2:
            feb_col = j
        if isinstance(v, str) and "1st nov 2024" in v.lower():
            nov_col = j
    if feb_col is None:
        raise ValueError("TOU normaliser: could not identify the Feb 2025 column in counts block")
    if nov_col is None:
        raise ValueError("TOU normaliser: could not identify the Nov 2024 column in counts block")

    counts_rows: List[Dict[str, Any]] = []
    current_sector: Optional[str] = None
    for i in range(counts_header + 1, min(counts_header + 20, len(df))):
        row = df.iloc[i].tolist()
        sector = _cell_text(row[0]).strip()
        if sector:
            current_sector = sector
        asset = _cell_text(row[1]).strip()
        if not asset:
            continue
        # Stop on the next block ("Assets" + Import/Export)
        if asset.lower() == "assets" and any("import" in _cell_text(v).lower() for v in row):
            break

        def _num(x: Any) -> Optional[float]:
            try:
                if x is None or (isinstance(x, float) and pd.isna(x)):
                    return None
                return float(x)
            except Exception:
                return None

        counts_rows.append(
            {
                "sector": current_sector,
                "asset": asset,
                "count_nov_2024": _num(row[nov_col]),
                "count_feb_2025": _num(row[feb_col]),
            }
        )

    # ---- Block 2: tariff metadata (Part 4.1) ----
    tariff_header = _find_row(
        df,
        lambda r: len(r) >= 7 and _cell_text(r[0]).lower() == "tariff name" and _cell_text(r[1]).lower() == "sector",
    )
    if tariff_header is None:
        raise ValueError("TOU normaliser: could not locate the tariff metadata header row")

    tariffs: List[Dict[str, Any]] = []
    for i in range(tariff_header + 1, min(tariff_header + 30, len(df))):
        row = df.iloc[i].tolist()
        name = _cell_text(row[0])
        if not name:
            break
        tariffs.append(
            {
                "tariff_name": name,
                "sector": _cell_text(row[1]),
                "asset_type": _cell_text(row[2]),
                "tariff_type": _cell_text(row[3]),
                "peak_times": _cell_text(row[4]),
                "offpeak_times": _cell_text(row[5]),
                "meter_count": row[6],
                "description": _cell_text(row[8]) if len(row) > 8 else "",
            }
        )

    # ---- Block 3: total turn-up/down differences (Part 4.2) ----
    totals: List[Dict[str, Any]] = []
    for i in range(len(df)):
        row = df.iloc[i].tolist()
        if (
            _cell_text(row[0]).lower() == "tariff name"
            and _cell_text(row[1]).lower() == "comparison"
            and "turn down" in _cell_text(row[2]).lower()
            and "turn up" in _cell_text(row[3]).lower()
        ):
            for j in range(i + 1, min(i + 20, len(df))):
                r = df.iloc[j].tolist()
                name = _cell_text(r[0])
                if not name:
                    break
                comparison = _cell_text(r[1])
                if not baseline_matcher(comparison):
                    continue
                totals.append(
                    {
                        "tariff_name": name,
                        "comparison": comparison,
                        "turn_down_mwh": r[2],
                        "turn_up_mwh": r[3],
                    }
                )
            if totals:
                break

    # ---- Block 4: daily average busiest (Part 4.3) ----
    busiest: List[Dict[str, Any]] = []
    for i in range(len(df)):
        row = df.iloc[i].tolist()
        if (
            _cell_text(row[0]).lower() == "tariff name"
            and _cell_text(row[1]).lower() == "comparison"
            and "daily avg busiest" in _cell_text(row[2]).lower()
            and "turn down mw" in " ".join(_cell_text(v).lower() for v in row)
        ):
            for j in range(i + 1, min(i + 20, len(df))):
                r = df.iloc[j].tolist()
                name = _cell_text(r[0])
                if not name:
                    break
                comparison = _cell_text(r[1])
                if not baseline_matcher(comparison):
                    continue
                busiest.append(
                    {
                        "tariff_name": name,
                        "comparison": comparison,
                        "busiest_turn_down_mwh": r[2],
                        "busiest_turn_up_mwh": r[3],
                        "busiest_turn_down_mw": r[4],
                        "busiest_turn_up_mw": r[5],
                    }
                )
            if busiest:
                break

    return {
        "counts": counts_rows,
        "tariffs": tariffs,
        "totals": totals,
        "busiest": busiest,
    }


def _write_part2_counts(ws, counts: List[Dict[str, Any]], notes: List[str]) -> None:
    """
    Fill Part 2 (managed assets) table with the Feb 2025 asset counts.

    We map the Feb 2025 snapshot into both 'No. in portfolio' and 'No. used in winter 2024-25'
    because this return provides point-in-time counts rather than a distinct winter-active measure.
    """
    # Template layout (Excel 1-based column numbers)
    #   sector at col 13, asset at col 14, portfolio at col 16, winter at col 18
    sector_col = 13
    asset_col = 14
    portfolio_col = 16
    winter_col = 18

    def norm_sector(s: str) -> str:
        t = (s or "").strip().lower()
        if "domestic" in t:
            return "Domestic"
        if "i&c" in t or "industrial" in t or t == "ic":
            return "I&C"
        return s

    def norm_asset(s: str) -> str:
        return (s or "").strip().lower()

    lookup: Dict[Tuple[str, str], Optional[float]] = {}
    for r in counts:
        sec = norm_sector(str(r.get("sector") or "")).strip().lower()
        asset = norm_asset(str(r.get("asset") or ""))
        lookup[(sec, asset)] = r.get("count_feb_2025")

    filled = 0
    current_sector_label: Optional[str] = None
    for row in range(10, 80):
        sec = ws.cell(row=row, column=sector_col).value
        asset = ws.cell(row=row, column=asset_col).value
        if sec is not None and str(sec).strip():
            current_sector_label = norm_sector(str(sec))
        if current_sector_label is None or asset is None:
            continue
        key = (str(current_sector_label).strip().lower(), str(asset).strip().lower())
        if key not in lookup:
            continue
        val = lookup.get(key)
        if val is None:
            continue
        try:
            num = float(val)
        except Exception:
            continue
        ws.cell(row=row, column=portfolio_col).value = num
        ws.cell(row=row, column=winter_col).value = num
        filled += 1

    notes.append(f"Part 2: filled {filled} asset count rows (Feb 2025 snapshot).")


def _write_part4_tables(ws, blocks: Dict[str, Any], notes: List[str]) -> None:
    """Fill Part 4 ToU tables (4.1/4.2/4.3)."""
    col_tariff = 2   # B
    col_sector = 7   # G
    col_asset = 9    # I
    col_type = 12    # L
    col_peak = 15    # O
    col_offpeak = 16 # P
    col_customers = 18 # R

    # Clear placeholder/example rows for 4.1 then write tariffs
    start_row = 79
    max_rows = 10
    for r in range(start_row, start_row + max_rows):
        for c in [col_tariff, col_sector, col_asset, col_type, col_peak, col_offpeak, col_customers]:
            ws.cell(row=r, column=c).value = None

    tariffs = blocks.get("tariffs") or []
    for idx, t in enumerate(tariffs[:max_rows]):
        r = start_row + idx
        ws.cell(row=r, column=col_tariff).value = t.get("tariff_name") or ""
        ws.cell(row=r, column=col_sector).value = t.get("sector") or ""
        ws.cell(row=r, column=col_asset).value = t.get("asset_type") or ""
        ws.cell(row=r, column=col_type).value = t.get("tariff_type") or ""
        ws.cell(row=r, column=col_peak).value = t.get("peak_times") or ""
        ws.cell(row=r, column=col_offpeak).value = t.get("offpeak_times") or ""
        ws.cell(row=r, column=col_customers).value = t.get("meter_count")

    notes.append(f"Part 4.1: wrote {min(len(tariffs), max_rows)} tariff metadata rows.")

    # 4.2 totals
    col_td = 7  # G
    col_tu = 9  # I
    start_row_42 = 97
    max_rows_42 = 10
    for r in range(start_row_42, start_row_42 + max_rows_42):
        for c in [col_tariff, col_td, col_tu]:
            ws.cell(row=r, column=c).value = None

    totals = blocks.get("totals") or []
    for idx, t in enumerate(totals[:max_rows_42]):
        r = start_row_42 + idx
        ws.cell(row=r, column=col_tariff).value = t.get("tariff_name") or ""
        ws.cell(row=r, column=col_td).value = t.get("turn_down_mwh")
        ws.cell(row=r, column=col_tu).value = t.get("turn_up_mwh")

    notes.append(
        "Part 4.2: totals mapped for selected comparison baseline." if totals else "Part 4.2: no totals block found for selected baseline."
    )

    # 4.3 daily busiest
    col_bmwh_down = 7   # G
    col_bmwh_up = 9     # I
    col_bmw_down = 12   # L
    col_bmw_up = 15     # O
    start_row_43 = 116
    max_rows_43 = 10
    for r in range(start_row_43, start_row_43 + max_rows_43):
        for c in [col_tariff, col_bmwh_down, col_bmwh_up, col_bmw_down, col_bmw_up]:
            ws.cell(row=r, column=c).value = None

    busiest = blocks.get("busiest") or []
    for idx, t in enumerate(busiest[:max_rows_43]):
        r = start_row_43 + idx
        ws.cell(row=r, column=col_tariff).value = t.get("tariff_name") or ""
        ws.cell(row=r, column=col_bmwh_down).value = t.get("busiest_turn_down_mwh")
        ws.cell(row=r, column=col_bmwh_up).value = t.get("busiest_turn_up_mwh")
        ws.cell(row=r, column=col_bmw_down).value = t.get("busiest_turn_down_mw")
        ws.cell(row=r, column=col_bmw_up).value = t.get("busiest_turn_up_mw")

    notes.append(
        "Part 4.3: daily busiest HH values mapped for selected baseline." if busiest else "Part 4.3: no daily busiest block found for selected baseline."
    )


def normalise_tou_return_to_v1_2(
    *,
    raw_path: Path,
    template_path: Path,
    output_path: Path,
    baseline_matcher: Callable[[str], bool] = _default_baseline_matcher,
) -> TOUReturnNormalisationResult:
    """Create a V1.2-shaped workbook populated from a non-template TOU submission."""
    raw_path = Path(raw_path)
    template_path = Path(template_path)
    output_path = Path(output_path)

    if not raw_path.exists():
        raise FileNotFoundError(f"Raw file not found: {raw_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"Template V1.2 file not found: {template_path}")

    blocks = _extract_tou_blocks(raw_path, baseline_matcher=baseline_matcher)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    wb = load_workbook(output_path)
    ws = wb.active

    notes: List[str] = []

    # Part 1: basic identifiers (best-effort)
    # Use a generic label for Company name to avoid writing provider identifiers into the normalised file.
    ws.cell(row=4, column=10).value = "Non-template TOU return (normalised)"
    ws.cell(row=13, column=10).value = "Supplier (non-template TOU return)"
    notes.append("Part 1: set Company name (generic label) and Company type.")

    _write_part2_counts(ws, blocks.get("counts") or [], notes)
    _write_part4_tables(ws, blocks, notes)

    ws.cell(row=72, column=2).value = (
        "NOTE (ADE): This file was normalised from a non-template TOU return into the V1.2 template layout. "
        "Part 4 values use a selected comparison baseline (default: any baseline containing 'flexible')."
    )

    wb.save(output_path)
    return TOUReturnNormalisationResult(output_path=output_path, notes=notes)

